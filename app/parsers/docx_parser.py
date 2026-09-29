import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.dml import MSO_COLOR_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.models.document import (
    ImageData,
    ParagraphData,
    ParsedDocxDocument,
    ParagraphType,
    SectionData,
    TableData,
    TextRunData,
)
from app.services.paragraph_classifier import classify_paragraph

def _length_to_cm(value) -> float | None:
    """Преобразует длину python-docx в сантиметры."""
    if value is None:
        return None
    return round(value.cm, 3)


def _length_to_pt(value) -> float | None:
    """Преобразует длину python-docx в пункты."""
    if value is None:
        return None
    return round(value.pt, 3)


def _alignment_to_str(alignment) -> str | None:
    """Преобразует тип выравнивания в понятную строку."""
    if alignment is None:
        return None

    mapping = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }
    return mapping.get(alignment, str(alignment))


def _get_page_orientation(section) -> str | None:
    if section.page_width is None or section.page_height is None:
        return None

    if section.page_width > section.page_height:
        return "landscape"

    return "portrait"


PAGE_FORMATS_CM = {
    "A0": (84.1, 118.9),
    "A1": (59.4, 84.1),
    "A2": (42.0, 59.4),
    "A3": (29.7, 42.0),
    "A4": (21.0, 29.7),
    "A5": (14.8, 21.0),
    "A6": (10.5, 14.8),
}


def _is_close(value: float | None, expected: float, tolerance: float = 0.2) -> bool:
    if value is None:
        return False

    return abs(value - expected) <= tolerance


def _get_page_format(section) -> str | None:
    width = _length_to_cm(section.page_width)
    height = _length_to_cm(section.page_height)

    if width is None or height is None:
        return None

    for format_name, (format_width, format_height) in PAGE_FORMATS_CM.items():
        if (
            _is_close(width, format_width)
            and _is_close(height, format_height)
        ) or (
            _is_close(width, format_height)
            and _is_close(height, format_width)
        ):
            return format_name

    return f"{width:g} x {height:g} cm"


def _get_part_text(part) -> str:
    return "\n".join(
        paragraph.text.strip()
        for paragraph in part.paragraphs
        if paragraph.text.strip()
    )


def _paragraph_has_page_field(paragraph: Paragraph) -> bool:
    for instr_text in paragraph._element.iter(qn("w:instrText")):
        if "PAGE" in (instr_text.text or "").upper():
            return True

    for fld_simple in paragraph._element.iter(qn("w:fldSimple")):
        if "PAGE" in (fld_simple.get(qn("w:instr")) or "").upper():
            return True

    return False


def _element_has_page_field(element) -> bool:
    for instr_text in element.iter(qn("w:instrText")):
        if "PAGE" in (instr_text.text or "").upper():
            return True

    for fld_simple in element.iter(qn("w:fldSimple")):
        if "PAGE" in (fld_simple.get(qn("w:instr")) or "").upper():
            return True

    return False


def _part_has_page_number(part) -> bool:
    if any(_paragraph_has_page_field(paragraph) for paragraph in part.paragraphs):
        return True

    # В реальных DOCX Word нередко помещает номер страницы внутрь w:sdt
    # (контент-контрола для колонтитула). Такие абзацы не всегда попадают
    # в python-docx part.paragraphs, поэтому делаем прямую проверку XML части.
    return _element_has_page_field(part._element)


def _get_paragraph_alignment(paragraph: Paragraph) -> str:
    if paragraph.alignment is not None:
        return _alignment_to_str(paragraph.alignment) or "left"

    if paragraph.style is not None:
        style_alignment = _get_style_alignment(paragraph.style)
        if style_alignment is not None:
            return style_alignment

    return "left"


def _get_page_number_alignment(part) -> str | None:
    for paragraph in part.paragraphs:
        if _paragraph_has_page_field(paragraph):
            return _get_paragraph_alignment(paragraph)

    for paragraph_element in part._element.iter(qn("w:p")):
        if not _element_has_page_field(paragraph_element):
            continue

        p_pr = paragraph_element.find(qn("w:pPr"))
        if p_pr is None:
            return "left"

        jc = p_pr.find(qn("w:jc"))
        if jc is None:
            return "left"

        val = jc.get(qn("w:val"))
        if val in {"left", "center", "right", "both", "justify"}:
            return "justify" if val == "both" else val
        return val or "left"

    return None


def _get_page_number_start(section) -> int | None:
    pg_num_type = section._sectPr.find(qn("w:pgNumType"))

    if pg_num_type is None:
        return None

    start = pg_num_type.get(qn("w:start"))

    if start is None:
        return None

    try:
        return int(start)
    except ValueError:
        return None


def _has_different_first_page_header_footer(section) -> bool:
    return section._sectPr.find(qn("w:titlePg")) is not None


TITLE_PAGE_KEYWORDS = {
    "министерство",
    "университет",
    "институт",
    "факультет",
    "кафедра",
    "выпускная квалификационная работа",
    "курсовая работа",
    "лабораторная работа",
    "отчет",
    "отчёт",
    "реферат",
    "тема",
    "выполнил",
    "выполнила",
    "студент",
    "обучающийся",
    "руководитель",
    "преподаватель",
    "проверил",
    "проверила",
}

TITLE_PAGE_STOP_WORDS = {
    "содержание",
    "оглавление",
    "введение",
    "заключение",
}


def _normalize_text_for_title_page(text: str) -> str:
    return " ".join(text.strip().lower().split())


def _looks_like_title_page_end(text: str) -> bool:
    normalized = _normalize_text_for_title_page(text).rstrip(".")

    if normalized in TITLE_PAGE_STOP_WORDS:
        return True

    return bool(re.match(r"^\d+(?:\.\d+)*\s+\S+", normalized))


def _title_page_score(paragraphs: list[ParagraphData]) -> int:
    joined_text = "\n".join(
        _normalize_text_for_title_page(paragraph.text)
        for paragraph in paragraphs
        if paragraph.text.strip()
    )

    score = 0
    for keyword in TITLE_PAGE_KEYWORDS:
        if keyword in joined_text:
            score += 1

    if re.search(r"\b20\d{2}\b", joined_text):
        score += 1

    if re.search(r"\b(москва|санкт-петербург|спб|казань|новосибирск)\b", joined_text):
        score += 1

    return score


def _mark_title_page_paragraphs(paragraphs: list[ParagraphData]) -> None:
    candidate_indexes: list[int] = []

    for paragraph in paragraphs[:40]:
        if _looks_like_title_page_end(paragraph.text):
            break

        candidate_indexes.append(paragraph.index)

    candidates = [paragraphs[index] for index in candidate_indexes]

    if _title_page_score(candidates) < 3:
        return

    for paragraph in candidates:
        paragraph.paragraph_type = ParagraphType.TITLE_PAGE
        paragraph.heading_level = None


def _logical_section_from_heading(text: str) -> str | None:
    normalized = _normalize_text_for_title_page(text).rstrip(".")
    normalized = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", normalized).rstrip(".")

    if normalized in {"введение"}:
        return "introduction"

    if normalized in {"заключение", "вывод", "выводы"}:
        return "conclusion"

    if normalized in {
        "список литературы",
        "список источников",
        "список использованной литературы",
        "список использованных источников",
        "библиографический список",
        "библиография",
    }:
        return "references"

    if normalized in {"приложение", "приложения"} or re.match(
        r"^приложение\s+[а-яёa-z](?:\.|\s|$)",
        normalized,
        re.IGNORECASE,
    ):
        return "appendix"

    return "main_part"


def _assign_logical_sections(paragraphs: list[ParagraphData]) -> None:
    current_section = "other"

    for paragraph in paragraphs:
        if paragraph.paragraph_type == ParagraphType.TITLE_PAGE:
            paragraph.logical_section = "title_page"
            continue

        if paragraph.paragraph_type == ParagraphType.HEADING and paragraph.heading_level == 1:
            current_section = _logical_section_from_heading(paragraph.text) or "main_part"
            paragraph.logical_section = current_section
            continue

        if paragraph.paragraph_type in {ParagraphType.FOOTNOTE, ParagraphType.TABLE_CELL}:
            paragraph.logical_section = None
            continue

        paragraph.logical_section = current_section

def _get_style_alignment(style) -> str | None:
    """
    Ищет выравнивание в стиле и его базовых стилях.
    """
    current_style = style

    while current_style is not None:
        alignment = current_style.paragraph_format.alignment
        if alignment is not None:
            return _alignment_to_str(alignment)

        current_style = current_style.base_style

    return None


def _get_effective_alignment(paragraph: Paragraph, document: Document) -> str | None:
    """
    Определяет выравнивание абзаца с учётом прямого форматирования и стилей.

    Если выравнивание нигде явно не задано, считаем его левым,
    потому что для Word это стандартное поведение по умолчанию.
    """
    if paragraph.alignment is not None:
        return _alignment_to_str(paragraph.alignment)

    if paragraph.style is not None:
        style_alignment = _get_style_alignment(paragraph.style)
        if style_alignment is not None:
            return style_alignment

    normal_style = document.styles["Normal"]
    normal_alignment = _get_style_alignment(normal_style)
    if normal_alignment is not None:
        return normal_alignment

    return "left"

def _get_first_non_empty_run(paragraph: Paragraph) -> Run | None:
    """
    Возвращает первый непустой run абзаца.

    В DOCX один абзац может состоять из нескольких run с разным форматированием.
    Пока для базовой версии берём первый непустой run.
    """
    for run in paragraph.runs:
        if run.text.strip():
            return run
    return None


def _get_effective_font_family(paragraph: Paragraph, document: Document) -> str | None:
    """
    Определяет шрифт с учётом прямого форматирования и стилей.

    Приоритет:
    1. прямое форматирование первого непустого run;
    2. стиль абзаца;
    3. стиль Normal.
    """
    run = _get_first_non_empty_run(paragraph)

    if run is not None and run.font.name is not None:
        return run.font.name

    if paragraph.style is not None and paragraph.style.font.name is not None:
        return paragraph.style.font.name

    normal_style = document.styles["Normal"]
    if normal_style.font.name is not None:
        return normal_style.font.name

    return None


def _get_effective_font_size_pt(paragraph: Paragraph, document: Document) -> float | None:
    """
    Определяет размер шрифта с учётом прямого форматирования и стилей.
    """
    run = _get_first_non_empty_run(paragraph)

    if run is not None and run.font.size is not None:
        return _length_to_pt(run.font.size)

    if paragraph.style is not None and paragraph.style.font.size is not None:
        return _length_to_pt(paragraph.style.font.size)

    normal_style = document.styles["Normal"]
    if normal_style.font.size is not None:
        return _length_to_pt(normal_style.font.size)

    return None


def _get_effective_bold(paragraph: Paragraph, document: Document) -> bool | None:
    """
    Определяет жирность с учётом прямого форматирования и стилей.

    Приоритет:
    1. прямое форматирование первого непустого run;
    2. стиль абзаца;
    3. стиль Normal.
    """
    run = _get_first_non_empty_run(paragraph)

    if run is not None and run.font.bold is not None:
        return run.font.bold

    if paragraph.style is not None and paragraph.style.font.bold is not None:
        return paragraph.style.font.bold

    normal_style = document.styles["Normal"]
    if normal_style.font.bold is not None:
        return normal_style.font.bold

    return None


def _get_effective_run_bool(
    run: Run,
    paragraph: Paragraph,
    document: Document,
    attribute: str,
) -> bool | None:
    value = getattr(run.font, attribute, None)
    if value is not None:
        return value

    if run.style is not None and getattr(run.style.font, attribute, None) is not None:
        return getattr(run.style.font, attribute, None)

    if paragraph.style is not None and getattr(paragraph.style.font, attribute, None) is not None:
        return getattr(paragraph.style.font, attribute, None)

    normal_style = document.styles["Normal"]
    if getattr(normal_style.font, attribute, None) is not None:
        return getattr(normal_style.font, attribute, None)

    return None


def _get_effective_paragraph_bold(paragraph: Paragraph, document: Document) -> bool | None:
    non_empty_runs = [run for run in paragraph.runs if run.text.strip()]

    if not non_empty_runs:
        return None

    return all(
        _get_effective_run_bool(run, paragraph, document, "bold") is True
        for run in non_empty_runs
    )


def _get_effective_paragraph_italic(paragraph: Paragraph, document: Document) -> bool | None:
    non_empty_runs = [run for run in paragraph.runs if run.text.strip()]

    if not non_empty_runs:
        return None

    return all(
        _get_effective_run_bool(run, paragraph, document, "italic") is True
        for run in non_empty_runs
    )


def _get_effective_run_font_size_pt(
    run: Run,
    paragraph: Paragraph,
    document: Document,
) -> float | None:
    if run.font.size is not None:
        return _length_to_pt(run.font.size)

    if run.style is not None and run.style.font.size is not None:
        return _length_to_pt(run.style.font.size)

    if paragraph.style is not None and paragraph.style.font.size is not None:
        return _length_to_pt(paragraph.style.font.size)

    normal_style = document.styles["Normal"]
    if normal_style.font.size is not None:
        return _length_to_pt(normal_style.font.size)

    return None


def _get_run_color(run: Run) -> str | None:
    color = run.font.color
    if color is None:
        return None

    if color.type == MSO_COLOR_TYPE.AUTO:
        return "000000"

    if color.rgb is None:
        return None

    return str(color.rgb).upper()


def _get_run_character_spacing(run: Run) -> int | None:
    run_properties = run._r.rPr
    if run_properties is None:
        return None

    spacing = run_properties.find(qn("w:spacing"))
    if spacing is None:
        return None

    value = spacing.get(qn("w:val"))
    if value is None:
        return None

    try:
        return int(value)
    except ValueError:
        return None


def _count_nonbreaking_spaces(text: str) -> int:
    return text.count("\u00a0")


def _get_break_flags(paragraph: Paragraph) -> dict[str, bool]:
    has_line_break = False
    has_page_break = False
    has_column_break = False

    for br in paragraph._element.iter(qn("w:br")):
        break_type = br.get(qn("w:type"))
        if break_type == "page":
            has_page_break = True
        elif break_type == "column":
            has_column_break = True
        else:
            has_line_break = True

    return {
        "has_line_break": has_line_break,
        "has_page_break": has_page_break,
        "has_column_break": has_column_break,
        "has_section_break": paragraph._p.pPr is not None
        and paragraph._p.pPr.find(qn("w:sectPr")) is not None,
    }


EQUATION_NUMBER_PATTERN = re.compile(r"\(\d+(?:\.\d+)*\)\s*$")


def _get_equation_flags(paragraph: Paragraph) -> dict[str, object]:
    equation_count = sum(1 for _ in paragraph._element.iter(qn("m:oMath")))
    equation_count += sum(1 for _ in paragraph._element.iter(qn("m:oMathPara")))
    text = paragraph.text.strip()
    number_match = EQUATION_NUMBER_PATTERN.search(text)
    has_image = _get_paragraph_image_count(paragraph) > 0
    is_standalone = bool(
        equation_count
        and (
            len(re.sub(EQUATION_NUMBER_PATTERN, "", text).strip()) <= 8
            or number_match
        )
    )

    return {
        "has_equation": equation_count > 0,
        "equation_count": equation_count,
        "equation_number": number_match.group(0) if number_match else None,
        "is_standalone_equation": is_standalone,
        "has_formula_like_image": bool(has_image and number_match),
    }


PRIVATE_DASH_MARKERS = "\uf02d"
PRIVATE_BULLET_MARKERS = "\uf076\uf0b7"

LIST_MARKER_PATTERNS = (
    ("dash", re.compile(rf"^[—–{PRIVATE_DASH_MARKERS}]\s+\S")),
    ("hyphen", re.compile(r"^-\s+\S")),
    ("bullet", re.compile(rf"^[•▪▫‣{PRIVATE_BULLET_MARKERS}]\s+\S")),
    ("number_bracket", re.compile(r"^\d+\)\s+\S")),
    ("number_dot", re.compile(r"^\d+\.\s+\S")),
    ("letter_bracket", re.compile(r"^[а-я]\)\s+\S", re.IGNORECASE)),
    ("letter_dot", re.compile(r"^[а-я]\.\s+\S", re.IGNORECASE)),
)


def _normalize_list_marker(marker: str) -> str:
    if marker in PRIVATE_DASH_MARKERS:
        return "—"
    if marker in PRIVATE_BULLET_MARKERS:
        return "•"
    return marker


def _get_list_marker(text: str) -> tuple[str | None, str | None]:
    stripped = text.strip()
    for marker_type, pattern in LIST_MARKER_PATTERNS:
        match = pattern.match(stripped)
        if match:
            return _normalize_list_marker(stripped.split(maxsplit=1)[0]), marker_type
    return None, None


def _has_automatic_numbering(paragraph: Paragraph) -> bool:
    p_pr = paragraph._p.pPr
    return bool(p_pr is not None and p_pr.numPr is not None)


def _get_automatic_list_marker_type(paragraph: Paragraph) -> str | None:
    marker, marker_type = _get_automatic_list_marker_details(paragraph)
    return marker_type


def _marker_type_from_automatic_format(fmt: str | None, marker: str | None) -> str:
    marker_text = (marker or "").strip()
    if marker_text in {"—", "–", *PRIVATE_DASH_MARKERS}:
        return "dash"
    if marker_text == "-":
        return "hyphen"
    if marker_text in {"•", "▪", "▫", "‣", *PRIVATE_BULLET_MARKERS}:
        return "bullet"
    if re.search(r"%\d+\)", marker_text):
        return "number_bracket"
    if re.search(r"%\d+\.", marker_text):
        return "number_dot"
    if fmt == "bullet":
        return "bullet" if marker_text else "bullet"
    if fmt in {"lowerLetter", "upperLetter", "lowerRoman", "upperRoman"}:
        return "letter_dot"
    return "number_dot"


def _get_automatic_list_marker_details(paragraph: Paragraph) -> tuple[str | None, str | None]:
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.numPr is None or p_pr.numPr.numId is None:
        return None, None

    num_id = str(p_pr.numPr.numId.val)
    ilvl = str(p_pr.numPr.ilvl.val) if p_pr.numPr.ilvl is not None else "0"
    numbering_part = getattr(paragraph.part, "numbering_part", None)
    if numbering_part is None:
        return None, "number_dot"

    numbering_root = numbering_part.element
    abstract_num_id: str | None = None
    for num in numbering_root.iter(qn("w:num")):
        if num.get(qn("w:numId")) != num_id:
            continue
        abstract = num.find(qn("w:abstractNumId"))
        if abstract is not None:
            abstract_num_id = abstract.get(qn("w:val"))
        break

    if abstract_num_id is None:
        return None, "number_dot"

    for abstract_num in numbering_root.iter(qn("w:abstractNum")):
        if abstract_num.get(qn("w:abstractNumId")) != abstract_num_id:
            continue
        level = None
        for candidate in abstract_num.iter(qn("w:lvl")):
            if candidate.get(qn("w:ilvl")) == ilvl:
                level = candidate
                break
        if level is None:
            return None, "number_dot"
        num_fmt = level.find(qn("w:numFmt"))
        fmt = num_fmt.get(qn("w:val")) if num_fmt is not None else None
        lvl_text = level.find(qn("w:lvlText"))
        marker = lvl_text.get(qn("w:val")) if lvl_text is not None else None
        return marker, _marker_type_from_automatic_format(fmt, marker)

    return None, "number_dot"


def _body_table_contexts(
    document: Document,
    parsed_paragraphs: list[ParagraphData],
) -> dict[int, dict[str, int | None]]:
    """
    Возвращает соседние абзацы для таблиц в порядке XML тела документа.

    python-docx отдаёт document.tables и document.paragraphs отдельными списками.
    Если просто сопоставлять n-ю таблицу с n-м заголовком, можно ошибиться:
    в документе могут быть таблицы без заголовка, заголовки в титульной части
    или подписи, которые идут не строго по счёту. Поэтому используем реальный
    порядок элементов body.
    """

    continuation_pattern = re.compile(
        r"^\s*(?:Продолжение|Окончание)\s+таблицы\s+\d+(?:\.\d+)*\b",
        re.IGNORECASE,
    )

    paragraph_by_body_index = {
        body_index: paragraph
        for body_index, paragraph in enumerate(
            p for p in parsed_paragraphs if p.table_index is None
        )
    }
    contexts: dict[int, dict[str, int | None]] = {}
    paragraph_counter = 0
    table_counter = 0
    previous_paragraph_index: int | None = None
    previous_table_caption_index: int | None = None
    used_caption_indexes: set[int] = set()

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = paragraph_by_body_index.get(paragraph_counter)
            paragraph_counter += 1
            if paragraph is None:
                continue
            previous_paragraph_index = paragraph.index
            if (
                (paragraph.caption_kind == "table" and paragraph.text.strip())
                or continuation_pattern.match(paragraph.text.strip())
            ):
                previous_table_caption_index = paragraph.index
            continue

        if child.tag != qn("w:tbl"):
            continue

        title_index = None
        if (
            previous_table_caption_index is not None
            and previous_table_caption_index not in used_caption_indexes
        ):
            title_index = previous_table_caption_index
            used_caption_indexes.add(previous_table_caption_index)

        contexts[table_counter] = {
            "paragraph_before_index": previous_paragraph_index,
            "title_paragraph_index": title_index,
        }
        previous_table_caption_index = None
        table_counter += 1

    return contexts


def _get_paragraph_image_count(paragraph: Paragraph) -> int:
    return sum(1 for _ in paragraph._element.iter(qn("w:drawing")))


def _get_paragraph_image_rids(paragraph: Paragraph) -> list[str]:
    rids: list[str] = []
    for blip in paragraph._element.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid:
            rids.append(rid)
    return rids


def _get_image_pixels(blob: bytes) -> tuple[int | None, int | None]:
    try:
        from PIL import Image
    except Exception:
        return None, None

    try:
        with Image.open(BytesIO(blob)) as image:
            return image.size
    except Exception:
        return None, None


def _get_image_extension(partname: str) -> str | None:
    suffix = Path(str(partname)).suffix.lower().lstrip(".")
    return suffix or None


def _table_has_grid(table) -> bool | None:
    borders = table._tbl.tblPr.find(qn("w:tblBorders")) if table._tbl.tblPr is not None else None
    if borders is None:
        return None

    visible_borders = [
        border
        for border in borders
        if border.get(qn("w:val")) not in {None, "nil", "none"}
    ]
    return bool(visible_borders)


def _table_border_signature(table) -> tuple[str | None, float | None]:
    borders = table._tbl.tblPr.find(qn("w:tblBorders")) if table._tbl.tblPr is not None else None
    if borders is None:
        return None, None

    for border in borders:
        color = border.get(qn("w:color"))
        size = border.get(qn("w:sz"))
        if color or size:
            return color, float(size) / 8 if size and size.isdigit() else None

    return None, None


def _row_has_tbl_header(row) -> bool:
    tr_pr = row._tr.trPr
    return tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None


def _row_has_cant_split(row) -> bool:
    tr_pr = row._tr.trPr
    return tr_pr is not None and tr_pr.find(qn("w:cantSplit")) is not None


def _row_height_pt(row) -> float | None:
    tr_pr = row._tr.trPr
    if tr_pr is None:
        return None

    tr_height = tr_pr.find(qn("w:trHeight"))
    if tr_height is None:
        return None

    value = tr_height.get(qn("w:val"))
    if value is None or not value.isdigit():
        return None

    return round(int(value) / 20, 2)


def _get_table_alignment(table) -> str | None:
    tbl_pr = table._tbl.tblPr
    jc = tbl_pr.find(qn("w:jc")) if tbl_pr is not None else None
    if jc is None:
        return None

    value = jc.get(qn("w:val"))
    return {"left": "left", "center": "center", "right": "right"}.get(value, value)


def _parse_footnotes(file_path: Path, document: Document, start_index: int) -> tuple[list[ParagraphData], bool]:
    if not file_path.exists():
        return [], False

    try:
        with ZipFile(file_path) as archive:
            if "word/footnotes.xml" not in archive.namelist():
                return [], False
            xml_text = archive.read("word/footnotes.xml").decode("utf-8", errors="ignore")
    except Exception:
        return [], False

    footnote_texts = re.findall(r"<w:footnote\b[^>]*>(.*?)</w:footnote>", xml_text, re.DOTALL)
    parsed: list[ParagraphData] = []
    for offset, raw_footnote in enumerate(footnote_texts):
        if 'w:type="separator"' in raw_footnote or 'w:type="continuationSeparator"' in raw_footnote:
            continue
        texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", raw_footnote, re.DOTALL)
        text = "".join(re.sub(r"<.*?>", "", item) for item in texts).strip()
        if not text:
            continue
        parsed.append(
            ParagraphData(
                index=start_index + len(parsed),
                text=text,
                font_family=None,
                font_size_pt=None,
                paragraph_type=ParagraphType.FOOTNOTE,
            )
        )

    return parsed, True


def _parse_footnote_settings(file_path: Path) -> tuple[str | None, str | None]:
    if not file_path.exists():
        return None, None

    try:
        with ZipFile(file_path) as archive:
            if "word/settings.xml" not in archive.namelist():
                return None, None
            xml_text = archive.read("word/settings.xml").decode("utf-8", errors="ignore")
    except Exception:
        return None, None

    footnote_pr = re.search(r"<w:footnotePr\b[^>]*>(.*?)</w:footnotePr>", xml_text, re.DOTALL)
    if not footnote_pr:
        return None, None

    content = footnote_pr.group(1)
    num_fmt = re.search(r"<w:numFmt\b[^>]*w:val=\"([^\"]+)\"", content)
    num_restart = re.search(r"<w:numRestart\b[^>]*w:val=\"([^\"]+)\"", content)
    return (
        num_fmt.group(1) if num_fmt else None,
        num_restart.group(1) if num_restart else None,
    )


def _has_character_style(run: Run) -> bool:
    if run.style is None:
        return False

    style_name = (run.style.name or "").strip().lower()
    return style_name not in {"", "default paragraph font", "шрифт абзаца по умолчанию"}


def _has_direct_formatting(run: Run) -> bool:
    direct_values = (
        run.font.bold,
        run.font.italic,
        run.font.underline,
        run.font.strike,
        getattr(run.font, "double_strike", None),
        run.font.small_caps,
        run.font.all_caps,
        run.font.size,
        run.font.color.rgb if run.font.color is not None else None,
        _get_run_character_spacing(run),
    )

    return any(value is not None for value in direct_values)


def _parse_runs(paragraph: Paragraph, document: Document) -> list[TextRunData]:
    runs: list[TextRunData] = []

    for index, run in enumerate(paragraph.runs):
        nonbreaking_space_count = _count_nonbreaking_spaces(run.text)
        runs.append(
            TextRunData(
                index=index,
                text=run.text,
                style_name=run.style.name if run.style else None,
                bold=_get_effective_run_bool(run, paragraph, document, "bold"),
                italic=_get_effective_run_bool(run, paragraph, document, "italic"),
                underline=bool(run.font.underline),
                strike=_get_effective_run_bool(run, paragraph, document, "strike"),
                double_strike=_get_effective_run_bool(run, paragraph, document, "double_strike"),
                small_caps=_get_effective_run_bool(run, paragraph, document, "small_caps"),
                all_caps=_get_effective_run_bool(run, paragraph, document, "all_caps"),
                font_size_pt=_get_effective_run_font_size_pt(run, paragraph, document),
                font_color=_get_run_color(run),
                character_spacing=_get_run_character_spacing(run),
                subscript=bool(run.font.subscript),
                superscript=bool(run.font.superscript),
                has_nonbreaking_space=nonbreaking_space_count > 0,
                nonbreaking_space_count=nonbreaking_space_count,
                has_character_style=_has_character_style(run),
                has_direct_formatting=_has_direct_formatting(run),
            )
        )

    return runs


def _get_effective_first_line_indent_cm(
    paragraph: Paragraph,
    document: Document,
) -> float | None:
    """
    Определяет абзацный отступ с учётом прямого форматирования и стилей.

    Если отступ не задан ни на абзаце, ни в стиле абзаца, ни в Normal,
    возвращается 0.0 — это эффективное значение Word по умолчанию.
    Симметрично _get_effective_left_indent_cm и _get_effective_right_indent_cm.
    """
    if paragraph.paragraph_format.first_line_indent is not None:
        return _length_to_cm(paragraph.paragraph_format.first_line_indent)

    if (
        paragraph.style is not None
        and paragraph.style.paragraph_format.first_line_indent is not None
    ):
        return _length_to_cm(paragraph.style.paragraph_format.first_line_indent)

    normal_style = document.styles["Normal"]
    if normal_style.paragraph_format.first_line_indent is not None:
        return _length_to_cm(normal_style.paragraph_format.first_line_indent)

    return 0.0

def _get_effective_left_indent_cm(
    paragraph: Paragraph,
    document: Document,
) -> float | None:
    """
    Определяет левый отступ абзаца с учётом прямого форматирования и стилей.
    """
    if paragraph.paragraph_format.left_indent is not None:
        return _length_to_cm(paragraph.paragraph_format.left_indent)

    if (
        paragraph.style is not None
        and paragraph.style.paragraph_format.left_indent is not None
    ):
        return _length_to_cm(paragraph.style.paragraph_format.left_indent)

    normal_style = document.styles["Normal"]
    if normal_style.paragraph_format.left_indent is not None:
        return _length_to_cm(normal_style.paragraph_format.left_indent)

    return 0.0


def _get_effective_right_indent_cm(
    paragraph: Paragraph,
    document: Document,
) -> float | None:
    """
    Определяет правый отступ абзаца с учётом прямого форматирования и стилей.
    """
    if paragraph.paragraph_format.right_indent is not None:
        return _length_to_cm(paragraph.paragraph_format.right_indent)

    if (
        paragraph.style is not None
        and paragraph.style.paragraph_format.right_indent is not None
    ):
        return _length_to_cm(paragraph.style.paragraph_format.right_indent)

    normal_style = document.styles["Normal"]
    if normal_style.paragraph_format.right_indent is not None:
        return _length_to_cm(normal_style.paragraph_format.right_indent)

    return 0.0

def _get_effective_line_spacing(
    paragraph: Paragraph,
    document: Document,
) -> float | None:
    """
    Определяет межстрочный интервал с учётом прямого форматирования и стилей.

    Если интервал задан числом, например 1.5, возвращаем его как есть.
    Если он задан абсолютной величиной, пока возвращаем None,
    потому что это другой тип проверки.
    """
    spacing = paragraph.paragraph_format.line_spacing

    if spacing is None and paragraph.style is not None:
        spacing = paragraph.style.paragraph_format.line_spacing

    if spacing is None:
        spacing = document.styles["Normal"].paragraph_format.line_spacing

    if isinstance(spacing, (int, float)):
        return float(spacing)

    return None


def _get_effective_space_before_pt(paragraph: Paragraph, document: Document) -> float | None:
    value = paragraph.paragraph_format.space_before

    if value is None and paragraph.style is not None:
        value = paragraph.style.paragraph_format.space_before

    if value is None:
        value = document.styles["Normal"].paragraph_format.space_before

    return _length_to_pt(value)


def _get_effective_space_after_pt(paragraph: Paragraph, document: Document) -> float | None:
    value = paragraph.paragraph_format.space_after

    if value is None and paragraph.style is not None:
        value = paragraph.style.paragraph_format.space_after

    if value is None:
        value = document.styles["Normal"].paragraph_format.space_after

    return _length_to_pt(value)


def _get_effective_paragraph_bool_format(
    paragraph: Paragraph,
    document: Document,
    attribute: str,
) -> bool | None:
    value = getattr(paragraph.paragraph_format, attribute)
    if value is not None:
        return value

    if paragraph.style is not None:
        value = getattr(paragraph.style.paragraph_format, attribute)
        if value is not None:
            return value

    return getattr(document.styles["Normal"].paragraph_format, attribute)


class DocxParser:
    """Парсер DOCX-документов."""

    def parse(self, file_path: str | Path) -> ParsedDocxDocument:
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"Файл не найден: {file_path}")

        if file_path.suffix.lower() != ".docx":
            raise ValueError("DocxParser поддерживает только файлы формата .docx")

        document = Document(file_path)

        parsed_paragraphs: list[ParagraphData] = []
        image_rid_to_paragraph_index: dict[str, int] = {}
        for index, paragraph in enumerate(document.paragraphs):
            break_flags = _get_break_flags(paragraph)
            equation_flags = _get_equation_flags(paragraph)
            list_marker, list_marker_type = _get_list_marker(paragraph.text)
            is_automatic_list = _has_automatic_numbering(paragraph)
            if is_automatic_list and list_marker_type is None:
                list_marker, list_marker_type = _get_automatic_list_marker_details(paragraph)
            image_rids = _get_paragraph_image_rids(paragraph)
            image_count = _get_paragraph_image_count(paragraph)
            for rid in image_rids:
                image_rid_to_paragraph_index.setdefault(rid, index)

            paragraph_data = ParagraphData(
                index=index,
                text=paragraph.text,
                style_name=paragraph.style.name if paragraph.style else None,
                font_family=_get_effective_font_family(paragraph, document),
                font_size_pt=_get_effective_font_size_pt(paragraph, document),
                bold=_get_effective_paragraph_bold(paragraph, document),
                italic=_get_effective_paragraph_italic(paragraph, document),
                runs=_parse_runs(paragraph, document),
                alignment=_get_effective_alignment(paragraph, document),
                first_line_indent_cm=_get_effective_first_line_indent_cm(
                    paragraph,
                    document,
                ),
                left_indent_cm=_get_effective_left_indent_cm(
                    paragraph,
                    document,
                ),
                right_indent_cm=_get_effective_right_indent_cm(
                    paragraph,
                    document,
                ),
                line_spacing=_get_effective_line_spacing(paragraph, document),
                space_before_pt=_get_effective_space_before_pt(paragraph, document),
                space_after_pt=_get_effective_space_after_pt(paragraph, document),
                keep_with_next=_get_effective_paragraph_bool_format(
                    paragraph,
                    document,
                    "keep_with_next",
                ),
                keep_together=_get_effective_paragraph_bool_format(
                    paragraph,
                    document,
                    "keep_together",
                ),
                page_break_before=_get_effective_paragraph_bool_format(
                    paragraph,
                    document,
                    "page_break_before",
                ),
                has_line_break=break_flags["has_line_break"],
                has_page_break=break_flags["has_page_break"],
                has_column_break=break_flags["has_column_break"],
                has_section_break=break_flags["has_section_break"],
                has_inline_image=image_count > 0,
                image_count=image_count,
                list_marker=list_marker,
                list_marker_type=list_marker_type,
                is_automatic_list=is_automatic_list,
                has_equation=bool(equation_flags["has_equation"]),
                equation_count=int(equation_flags["equation_count"]),
                has_formula_like_image=bool(equation_flags["has_formula_like_image"]),
                is_standalone_equation=bool(equation_flags["is_standalone_equation"]),
                equation_number=equation_flags["equation_number"],
            )

            parsed_paragraphs.append(classify_paragraph(paragraph_data))

        _mark_title_page_paragraphs(parsed_paragraphs)

        parsed_images: list[ImageData] = []
        for image_index, shape in enumerate(document.inline_shapes):
            rid = None
            for blip in shape._inline.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                if rid:
                    break

            image_part = document.part.related_parts.get(rid) if rid else None
            width_px = height_px = None
            extension = None
            if image_part is not None:
                width_px, height_px = _get_image_pixels(image_part.blob)
                extension = _get_image_extension(image_part.partname)

            parsed_images.append(
                ImageData(
                    index=image_index,
                    paragraph_index=image_rid_to_paragraph_index.get(rid),
                    width_cm=_length_to_cm(shape.width),
                    height_cm=_length_to_cm(shape.height),
                    width_px=width_px,
                    height_px=height_px,
                    extension=extension,
                    is_inline=True,
                )
            )

        parsed_tables: list[TableData] = []
        next_paragraph_index = len(parsed_paragraphs)
        table_contexts = _body_table_contexts(document, parsed_paragraphs)

        for table_index, table in enumerate(document.tables):
            cell_paragraphs: list[ParagraphData] = []
            empty_rows: list[int] = []
            enlarged_rows: list[int] = []
            has_cant_split_rows = False
            repeats_header = False

            for row_index, row in enumerate(table.rows):
                row_text = " ".join(cell.text.strip() for cell in row.cells).strip()
                if not row_text:
                    empty_rows.append(row_index)

                row_height = _row_height_pt(row)
                if row_height is not None and row_height > 30:
                    enlarged_rows.append(row_index)

                has_cant_split_rows = has_cant_split_rows or _row_has_cant_split(row)
                repeats_header = repeats_header or _row_has_tbl_header(row)

                seen_cells: set[int] = set()
                for column_index, cell in enumerate(row.cells):
                    cell_id = id(cell._tc)
                    if cell_id in seen_cells:
                        continue
                    seen_cells.add(cell_id)

                    for cell_paragraph in cell.paragraphs:
                        break_flags = _get_break_flags(cell_paragraph)
                        equation_flags = _get_equation_flags(cell_paragraph)
                        list_marker, list_marker_type = _get_list_marker(cell_paragraph.text)
                        is_automatic_list = _has_automatic_numbering(cell_paragraph)
                        if is_automatic_list and list_marker_type is None:
                            list_marker, list_marker_type = _get_automatic_list_marker_details(cell_paragraph)
                        paragraph_data = ParagraphData(
                            index=next_paragraph_index,
                            text=cell_paragraph.text,
                            style_name=cell_paragraph.style.name if cell_paragraph.style else None,
                            font_family=_get_effective_font_family(cell_paragraph, document),
                            font_size_pt=_get_effective_font_size_pt(cell_paragraph, document),
                            bold=_get_effective_paragraph_bold(cell_paragraph, document),
                            italic=_get_effective_paragraph_italic(cell_paragraph, document),
                            runs=_parse_runs(cell_paragraph, document),
                            alignment=_get_effective_alignment(cell_paragraph, document),
                            first_line_indent_cm=_get_effective_first_line_indent_cm(
                                cell_paragraph,
                                document,
                            ),
                            left_indent_cm=_get_effective_left_indent_cm(
                                cell_paragraph,
                                document,
                            ),
                            right_indent_cm=_get_effective_right_indent_cm(
                                cell_paragraph,
                                document,
                            ),
                            line_spacing=_get_effective_line_spacing(cell_paragraph, document),
                            space_before_pt=_get_effective_space_before_pt(cell_paragraph, document),
                            space_after_pt=_get_effective_space_after_pt(cell_paragraph, document),
                            keep_with_next=_get_effective_paragraph_bool_format(
                                cell_paragraph,
                                document,
                                "keep_with_next",
                            ),
                            keep_together=_get_effective_paragraph_bool_format(
                                cell_paragraph,
                                document,
                                "keep_together",
                            ),
                            has_line_break=break_flags["has_line_break"],
                            has_page_break=break_flags["has_page_break"],
                            has_column_break=break_flags["has_column_break"],
                            has_section_break=break_flags["has_section_break"],
                            list_marker=list_marker,
                            list_marker_type=list_marker_type,
                            is_automatic_list=is_automatic_list,
                            has_equation=bool(equation_flags["has_equation"]),
                            equation_count=int(equation_flags["equation_count"]),
                            has_formula_like_image=bool(equation_flags["has_formula_like_image"]),
                            is_standalone_equation=bool(equation_flags["is_standalone_equation"]),
                            equation_number=equation_flags["equation_number"],
                            paragraph_type=ParagraphType.TABLE_CELL,
                            table_index=table_index,
                            table_row_index=row_index,
                            table_column_index=column_index,
                        )
                        cell_paragraphs.append(paragraph_data)
                        parsed_paragraphs.append(paragraph_data)
                        next_paragraph_index += 1

            border_color, border_size = _table_border_signature(table)
            table_context = table_contexts.get(table_index, {})
            title_index = table_context.get("title_paragraph_index")
            parsed_tables.append(
                TableData(
                    index=table_index,
                    rows_count=len(table.rows),
                    columns_count=len(table.columns),
                    paragraph_before_index=table_context.get("paragraph_before_index"),
                    title_paragraph_index=title_index,
                    has_header_row=bool(table.rows and any(cell.text.strip() for cell in table.rows[0].cells)),
                    alignment=_get_table_alignment(table),
                    has_grid=_table_has_grid(table),
                    border_color=border_color,
                    border_size=border_size,
                    has_cant_split_rows=has_cant_split_rows,
                    repeats_header=repeats_header,
                    empty_row_indexes=empty_rows,
                    enlarged_row_indexes=enlarged_rows,
                    cell_paragraphs=cell_paragraphs,
                )
            )

        footnotes, has_footnotes_part = _parse_footnotes(
            file_path,
            document,
            next_paragraph_index,
        )
        footnote_number_format, footnote_number_restart = _parse_footnote_settings(file_path)
        parsed_paragraphs.extend(footnotes)

        _assign_logical_sections(parsed_paragraphs)

        parsed_sections: list[SectionData] = []
        for index, section in enumerate(document.sections):
            header_text = _get_part_text(section.header)
            footer_text = _get_part_text(section.footer)
            has_header_page_number = _part_has_page_number(section.header)
            has_footer_page_number = _part_has_page_number(section.footer)
            has_first_page_header_page_number = _part_has_page_number(section.first_page_header)
            has_first_page_footer_page_number = _part_has_page_number(section.first_page_footer)
            has_first_page_page_number = (
                has_first_page_header_page_number
                or has_first_page_footer_page_number
            )
            has_page_number = has_header_page_number or has_footer_page_number
            page_number_start = _get_page_number_start(section)
            page_number_area = None
            page_number_alignment = None
            first_page_page_number_area = None

            if has_footer_page_number:
                page_number_area = "footer"
                page_number_alignment = _get_page_number_alignment(section.footer)
            elif has_header_page_number:
                page_number_area = "header"
                page_number_alignment = _get_page_number_alignment(section.header)

            if has_first_page_footer_page_number:
                first_page_page_number_area = "footer"
            elif has_first_page_header_page_number:
                first_page_page_number_area = "header"

            parsed_sections.append(
                SectionData(
                    index=index,
                    page_width_cm=_length_to_cm(section.page_width),
                    page_height_cm=_length_to_cm(section.page_height),
                    page_format=_get_page_format(section),
                    page_orientation=_get_page_orientation(section),
                    top_margin_cm=_length_to_cm(section.top_margin),
                    bottom_margin_cm=_length_to_cm(section.bottom_margin),
                    left_margin_cm=_length_to_cm(section.left_margin),
                    right_margin_cm=_length_to_cm(section.right_margin),
                    header_distance_cm=_length_to_cm(section.header_distance),
                    footer_distance_cm=_length_to_cm(section.footer_distance),
                    header_text=header_text,
                    footer_text=footer_text,
                    has_header=bool(header_text),
                    has_footer=bool(footer_text) or has_footer_page_number,
                    has_page_number=has_page_number,
                    has_footer_page_number=has_footer_page_number,
                    has_first_page_page_number=has_first_page_page_number,
                    first_page_has_page_number=has_first_page_page_number,
                    has_first_page_footer_page_number=has_first_page_footer_page_number,
                    page_number_area=page_number_area,
                    page_number_alignment=page_number_alignment,
                    first_page_page_number_area=first_page_page_number_area,
                    page_number_start=page_number_start,
                    page_number_restart=page_number_start is not None and index > 0,
                    different_first_page_header_footer=(
                        _has_different_first_page_header_footer(section)
                    ),
                )
            )

        return ParsedDocxDocument(
            file_path=str(file_path),
            paragraphs=parsed_paragraphs,
            sections=parsed_sections,
            images=parsed_images,
            tables=parsed_tables,
            footnotes=footnotes,
            has_footnotes_part=has_footnotes_part,
            footnote_number_format=footnote_number_format,
            footnote_number_restart=footnote_number_restart,
        )
